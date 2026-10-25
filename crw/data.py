"""Bounded imports and deterministic, entirely fictional demonstration data."""
from __future__ import annotations
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from random import Random
import json
import zipfile
import pandas as pd

SCHEMA = {
    'patients': ['patient_id', 'sex'],
    'encounters': ['encounter_id', 'patient_id', 'age', 'admission_date', 'discharge_date', 'diagnosis', 'procedure', 'surgery_date'],
    'labs': ['result_id', 'encounter_id', 'analyte', 'value', 'unit', 'specimen', 'sample_time'],
}
ALIASES = {
    '患者编号': 'patient_id', '患者ID': 'patient_id', '性别': 'sex',
    '住院编号': 'encounter_id', '就诊编号': 'encounter_id', '年龄': 'age',
    '入院时间': 'admission_date', '出院时间': 'discharge_date', '诊断': 'diagnosis',
    '手术方式': 'procedure', '手术时间': 'surgery_date', '检验编号': 'result_id',
    '检验项目': 'analyte', '结果': 'value', '单位': 'unit', '标本': 'specimen', '采样时间': 'sample_time',
}
DIAGNOSES = {'OVCF': '骨质疏松性椎体压缩骨折', 'SPINAL_INFECTION': '脊柱感染', 'DEGENERATIVE': '退变性疾病', 'OTHER': '其他'}
PROCEDURES = {'PKP': 'PKP', 'PVP': 'PVP', 'CONSERVATIVE': '保守治疗', 'UNKNOWN': '未知'}
MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 100_000
OPTIONAL = {'patients': [], 'encounters': [], 'labs': ['source_report_id', 'source_file', 'source_json_path']}

@dataclass
class Bundle:
    tables: dict[str, pd.DataFrame]
    source_hashes: dict[str, str]
    label: str
    synthetic: bool = False
    source_warnings: list[dict] = field(default_factory=list)


def from_tables(tables: dict[str, pd.DataFrame], label='上传数据', synthetic=False) -> Bundle:
    copied, hashes = {}, {}
    for table, columns in SCHEMA.items():
        if table not in tables:
            raise ValueError(f'缺少 {table} 表。')
        frame = tables[table].copy(deep=True)
        if len(frame) > MAX_ROWS:
            raise ValueError(f'{table} 超过 {MAX_ROWS:,} 行限制。')
        frame.columns = [ALIASES.get(str(c).strip(), str(c).strip()) for c in frame.columns]
        if frame.columns.duplicated().any():
            raise ValueError(f'{table} 字段映射后重名，请删除重复列。')
        required = [c for c in columns if c != 'surgery_date']
        missing = sorted(set(required) - set(frame.columns))
        if missing:
            raise ValueError(f'{table} 缺少字段：{", ".join(missing)}')
        if 'surgery_date' in columns and 'surgery_date' not in frame:
            frame['surgery_date'] = ''
        hashes[table] = sha256(frame.to_csv(index=False).encode('utf-8')).hexdigest()
        selected = columns + [c for c in OPTIONAL[table] if c in frame]
        copied[table] = frame.loc[:, selected].reset_index(drop=True).astype('string').fillna('')
    return Bundle(copied, hashes, label, synthetic)


def read_csv_files(files: dict[str, bytes]) -> Bundle:
    tables = {}
    for name in SCHEMA:
        raw = files.get(name, b'')
        if not raw or len(raw) > MAX_BYTES:
            raise ValueError(f'{name} 为空或超过10 MB。')
        try:
            tables[name] = pd.read_csv(BytesIO(raw), dtype=str, keep_default_na=False, encoding='utf-8-sig')
        except Exception as exc:
            raise ValueError(f'{name} 无法读取，请使用 UTF-8 CSV。') from exc
    bundle = from_tables(tables)
    bundle.source_hashes = {name: sha256(raw).hexdigest() for name, raw in files.items() if name in SCHEMA}
    return bundle


def _sheet_table(sheet) -> pd.DataFrame:
    if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > 100:
        raise ValueError(f'{sheet.title} 的行列数超过限制。')
    sheet.reset_dimensions()
    values = []
    for index, row in enumerate(sheet.iter_rows()):
        if index > MAX_ROWS or len(row) > 100:
            raise ValueError(f'{sheet.title} 的行列数超过限制。')
        if any(c.data_type == 'f' for c in row):
            raise ValueError('导入表含公式，请先另存为值；工具不会执行或计算公式。')
        values.append([c.value if c.value is not None else '' for c in row])
    if not values:
        raise ValueError(f'{sheet.title} 为空。')
    return pd.DataFrame(values[1:], columns=values[0])


def read_workbook(raw: bytes) -> Bundle:
    if not raw or len(raw) > MAX_BYTES:
        raise ValueError('工作簿为空或超过10 MB。')
    try:
        with zipfile.ZipFile(BytesIO(raw)) as archive:
            items = archive.infolist()
            if len(items) > 1000 or sum(i.file_size for i in items) > 50 * 1024 * 1024:
                raise ValueError('工作簿展开后过大。')
            if any(i.file_size > max(i.compress_size, 1) * 200 for i in items):
                raise ValueError('工作簿压缩比异常。')
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            legacy = {'病例总表', '检验项目明细'}.issubset(workbook.sheetnames)
            if legacy:
                legacy_tables = {}
                for name in ['病例总表', '检验项目明细', '质量核对']:
                    if name not in workbook.sheetnames:
                        continue
                    legacy_tables[name] = _sheet_table(workbook[name])
                bundle = from_legacy_tables(legacy_tables)
                bundle.source_hashes['workbook'] = sha256(raw).hexdigest()
                return bundle
            tables = {}
            for name in SCHEMA:
                if name not in workbook.sheetnames:
                    raise ValueError(f'工作簿缺少 {name} 工作表。')
                tables[name] = _sheet_table(workbook[name])
        finally:
            workbook.close()
        bundle = from_tables(tables)
        bundle.source_hashes['workbook'] = sha256(raw).hexdigest()
        return bundle
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('无法读取工作簿，请使用包含 patients、encounters、labs 三张表的 .xlsx。') from exc


def from_legacy_tables(tables: dict[str, pd.DataFrame]) -> Bundle:
    """Adapt a known collector output schema without using names as identity keys."""
    cases, items = tables['病例总表'].copy(), tables['检验项目明细'].copy()
    case_fields = ['原患者ID', '住院关联键', '性别', '年龄（岁）', '入院时间（原试跑表）', '出院时间（原试跑表）', '原表诊断（待核定）']
    lab_fields = ['检验项目关联键', '住院关联键', '项目名称', '结果原文', '单位原文', '标本类型', '送检时间']
    for frame, required in [(cases, case_fields), (items, lab_fields)]:
        missing = set(required) - set(frame.columns)
        if missing:
            raise ValueError('旧格式缺少必要字段：' + ', '.join(sorted(missing)))
    warnings = [{'code': 'SOURCE_DIAGNOSIS_PENDING', 'count': len(cases), 'message': '按原表待核定诊断建立候选队列，不自动判断最终临床纳入。'}]
    patient_rows = []
    for pid, group in cases.groupby('原患者ID', dropna=False, sort=False):
        sexes = group['性别'].fillna('').astype(str).str.strip().unique()
        sex = sexes[0] if len(sexes) == 1 else 'U'
        if len(sexes) > 1:
            warnings.append({'code': 'SOURCE_SEX_CONFLICT', 'count': 1, 'message': '同一来源患者键的性别不一致，记为未知待核对。'})
        patient_rows.append({'patient_id': pid, 'sex': sex})
    encounters = cases.rename(columns={'原患者ID': 'patient_id', '住院关联键': 'encounter_id', '年龄（岁）': 'age',
        '入院时间（原试跑表）': 'admission_date', '出院时间（原试跑表）': 'discharge_date', '原表诊断（待核定）': 'diagnosis'})
    encounters['procedure'] = 'UNKNOWN'
    encounters['surgery_date'] = ''
    labs = items.rename(columns={'检验项目关联键': 'result_id', '住院关联键': 'encounter_id', '项目名称': 'analyte',
        '结果原文': 'value', '单位原文': 'unit', '标本类型': 'specimen', '送检时间': 'sample_time',
        '报告关联键': 'source_report_id', '源文件': 'source_file', '源JSON路径': 'source_json_path'})
    if '质量核对' in tables and '核对代码' in tables['质量核对']:
        for code, n in tables['质量核对']['核对代码'].fillna('SOURCE_UNKNOWN').value_counts().items():
            warnings.append({'code': str(code), 'count': int(n), 'message': '保留上游质量标记；结构关联成功不代表该问题已解决。'})
    bundle = from_tables({'patients': pd.DataFrame(patient_rows, columns=SCHEMA['patients']), 'encounters': encounters, 'labs': labs}, '已有中文采集整理工作簿')
    bundle.source_warnings = warnings
    return bundle


def make_demo(seed: int = 20261025) -> Bundle:
    rng = Random(seed)
    patients = [{'patient_id': f'SYN-P{i:03}', 'sex': rng.choice(['M', 'F'])} for i in range(1, 81)]
    encounters, labs = [], []
    diagnoses = list(DIAGNOSES)
    for i in range(1, 101):
        pid = f'SYN-P{((i - 1) % 80) + 1:03}'
        admit = datetime(2025, 1, 1) + timedelta(days=i * 4)
        discharge = admit + timedelta(days=rng.randint(3, 20))
        diagnosis = diagnoses[(i - 1) % len(diagnoses)]
        procedure = rng.choice(['PKP', 'PVP']) if diagnosis == 'OVCF' else rng.choice(['CONSERVATIVE', 'UNKNOWN'])
        encounters.append({'encounter_id': f'SYN-E{i:03}', 'patient_id': pid, 'age': str(rng.randint(20, 88)),
                           'admission_date': admit.isoformat(sep=' '), 'discharge_date': discharge.isoformat(sep=' '),
                           'diagnosis': diagnosis, 'procedure': procedure,
                           'surgery_date': (admit + timedelta(days=1)).isoformat(sep=' ') if procedure in ['PVP', 'PKP'] else ''})
        for j, (analyte, unit, value, day) in enumerate([
            ('CRP', 'mg/L', f'{rng.uniform(0.4, 120):.2f}', 0),
            ('CRP', 'mg/dL', f'{rng.uniform(0.04, 9):.2f}', 3),
            ('WBC', '10^9/L', f'{rng.uniform(3, 15):.2f}', 0),
            ('ESR', 'mm/h', str(rng.randint(3, 100)), 0),
        ]):
            labs.append({'result_id': f'SYN-L{i:03}-{j}', 'encounter_id': f'SYN-E{i:03}', 'analyte': analyte,
                         'value': value, 'unit': unit, 'specimen': 'whole_blood' if analyte in ['WBC', 'ESR'] else 'serum',
                         'sample_time': (admit + timedelta(days=day, hours=6)).isoformat(sep=' ')})
    # Deliberate, documented defects. No real records or identifiers enter this generator.
    patients[3]['sex'] = ''
    encounters[1]['age'] = ''
    encounters[2]['age'] = '145'
    encounters[3]['discharge_date'] = '2024-01-01'
    encounters[4]['patient_id'] = 'SYN-ORPHAN'
    encounters[5]['admission_date'] = 'not-a-date'
    encounters[6]['diagnosis'] = ''
    encounters.append(dict(encounters[9]))
    labs[0]['value'] = '<5'
    labs[4]['unit'] = 'unknown-unit'
    labs[8]['sample_time'] = '2024-01-01'
    labs[12]['value'] = '-2'
    labs[16]['analyte'] = 'hsCRP'
    labs[20]['value'] = ''
    labs[24]['specimen'] = 'csf'
    labs[28]['encounter_id'] = 'SYN-ORPHAN'
    conflict = dict(labs[40]); conflict['result_id'] = 'SYN-CONFLICT'; conflict['value'] = '999'
    labs.append(conflict)
    duplicate = dict(labs[44]); duplicate['result_id'] = 'SYN-REPEAT'; labs.append(duplicate)
    labs.append(dict(labs[48]))
    return from_tables({'patients': pd.DataFrame(patients), 'encounters': pd.DataFrame(encounters), 'labs': pd.DataFrame(labs)}, '合成演示数据', True)


def save_demo(folder: Path) -> None:
    bundle = make_demo()
    folder.mkdir(parents=True, exist_ok=True)
    for name, frame in bundle.tables.items():
        frame.to_csv(folder / f'{name}.csv', index=False, encoding='utf-8-sig')
    (folder / 'manifest.json').write_text(json.dumps({'synthetic': True, 'seed': 20261025,
        'rows': {name: len(frame) for name, frame in bundle.tables.items()}, 'source_hashes': bundle.source_hashes}, ensure_ascii=False, indent=2))
