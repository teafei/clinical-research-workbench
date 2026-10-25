"""Clinical data edge cases; all fixtures are synthetic."""
import unittest
from io import BytesIO
import zipfile
import json
import pandas as pd
from openpyxl import Workbook, load_workbook

from crw.data import from_tables, make_demo, read_workbook, read_csv_files, from_legacy_tables
from crw.pipeline import prepare, build_cohort, CohortSpec, descriptive_tables, audit_manifest
from crw.export import export_package, workbook_bytes, csv_bytes


def small_bundle(labs=None, encounters=None):
    patients = pd.DataFrame([{'patient_id': '0001', 'sex': 'F'}, {'patient_id': '0002', 'sex': 'M'}])
    episodes = encounters if encounters is not None else [
        dict(encounter_id='E1', patient_id='0001', age='60', admission_date='2025-01-01 00:00:00',
             discharge_date='2025-01-10 00:00:00', diagnosis='OVCF', procedure='PVP'),
        dict(encounter_id='E2', patient_id='0001', age='61', admission_date='2025-02-01 23:59:59',
             discharge_date='2025-02-10 00:00:00', diagnosis='OVCF', procedure='PKP'),
        dict(encounter_id='E3', patient_id='0002', age='50', admission_date='2025-02-02 00:00:00',
             discharge_date='', diagnosis='SPINAL_INFECTION', procedure='UNKNOWN')]
    if labs is None:
        labs = [lab('R1')]
    return from_tables({'patients': patients, 'encounters': pd.DataFrame(episodes), 'labs': pd.DataFrame(labs)})


def lab(rid='R1', **kwargs):
    return dict(result_id=rid, encounter_id='E1', analyte='CRP', value='4', unit='mg/L', specimen='serum',
                sample_time='2025-01-01 01:00:00') | kwargs


class PipelineTests(unittest.TestCase):
    def test_demo_is_deterministic_and_inputs_unchanged(self):
        b = make_demo()
        copies = {k: v.copy(deep=True) for k, v in b.tables.items()}
        p = prepare(b)
        c, f = build_cohort(p, CohortSpec())
        self.assertEqual((len(b.tables['patients']), len(b.tables['encounters']), len(b.tables['labs'])), (80, 101, 403))
        self.assertEqual((len(p.usable_encounters), len(c), c.patient_id.nunique()), (94, 80, 80))
        self.assertEqual(b.source_hashes, make_demo().source_hashes)
        for k in copies:
            pd.testing.assert_frame_equal(copies[k], b.tables[k])
        self.assertTrue((f.before - f.after == f.excluded).all())

    def test_duplicate_primary_keys_and_orphan_links_block_whole_groups(self):
        b = small_bundle()
        b.tables['patients'] = pd.concat([b.tables['patients'], b.tables['patients'].iloc[[0]]], ignore_index=True)
        b.tables['encounters'].loc[2, 'patient_id'] = 'MISSING'
        p = prepare(b)
        self.assertTrue(p.usable_encounters.empty)
        self.assertTrue(p.usable_labs.empty)
        self.assertEqual(p.issues.code.eq('DUPLICATE_KEY').sum(), 2)

    def test_conversion_and_baseline_separate_censoring_specimen_and_analyte(self):
        records = [lab('C', value='<5', sample_time='2025-01-01 00:01:00'),
                   lab('H', analyte='hsCRP', value='99', sample_time='2025-01-01 00:02:00'),
                   lab('S', specimen='csf', value='888', sample_time='2025-01-01 00:03:00'),
                   lab('P', specimen='plasma', value='777', sample_time='2025-01-01 00:04:00'),
                   lab('B', value='1.2', unit='mg/dL')]
        p = prepare(small_bundle(records))
        c, _ = build_cohort(p, CohortSpec())
        self.assertEqual(c.loc[c.encounter_id.eq('E1'), 'crp_first_mg_L'].item(), 12)
        censored = p.usable_labs.set_index('result_id').loc['C']
        self.assertEqual(censored.comparator, '<')
        self.assertEqual(censored.numeric_value, 5)
        self.assertEqual(p.usable_labs.set_index('result_id').loc['H'].analyte, 'HSCRP')

    def test_unsupported_units_and_negative_known_values_are_not_usable(self):
        p = prepare(small_bundle([lab('U', unit='mg/unknown'), lab('N', value='-1')]))
        self.assertTrue(p.usable_labs.empty)
        self.assertIn('UNIT_UNSUPPORTED', p.issues.code.tolist())
        c, _ = build_cohort(p, CohortSpec())
        self.assertTrue(c.crp_first_mg_L.isna().all())

    def test_window_inclusive_and_discharge_boundary(self):
        p = prepare(small_bundle([lab('PRE', sample_time='2024-12-31 23:59:59'),
                                  lab('EDGE', sample_time='2025-01-04 00:00:00', value='9'),
                                  lab('LATE', sample_time='2025-01-04 00:00:01', value='99'),
                                  lab('POST', sample_time='2025-01-11', value='999')]))
        c, _ = build_cohort(p, CohortSpec(baseline_hours=72))
        self.assertEqual(c.set_index('encounter_id').loc['E1', 'crp_first_mg_L'], 9)
        self.assertEqual(p.issues.code.eq('OUTSIDE_ENCOUNTER').sum(), 2)

    def test_conflicts_block_whole_measurement_group_and_repeats_keep_sources(self):
        p = prepare(small_bundle([lab('A'), lab('B', value='5'), lab('C', sample_time='2025-01-02', value='6'),
                                  lab('D', sample_time='2025-01-02', value='6')]))
        self.assertEqual(len(p.usable_labs), 1)
        self.assertEqual(p.usable_labs.iloc[0].source_result_ids, 'C|D')
        self.assertEqual(p.issues.code.eq('MEASUREMENT_CONFLICT').sum(), 2)
        self.assertEqual(p.issues.code.eq('MEASUREMENT_REPEAT').sum(), 1)

    def test_text_results_not_misclassified_as_equivalent_missing_values(self):
        p = prepare(small_bundle([lab('A', analyte='CULTURE', value='阳性', unit='', specimen='blood'),
                                  lab('B', analyte='CULTURE', value='阴性', unit='', specimen='blood')]))
        self.assertTrue(p.usable_labs.empty)
        self.assertEqual(p.issues.code.eq('MEASUREMENT_CONFLICT').sum(), 2)

    def test_first_admission_is_selected_after_filters_and_end_date_includes_day(self):
        p = prepare(small_bundle())
        c, _ = build_cohort(p, CohortSpec(procedures=('PKP',), end_date='2025-02-01'))
        self.assertEqual(c.encounter_id.tolist(), ['E2'])
        c, _ = build_cohort(p, CohortSpec(first_per_patient=False))
        self.assertEqual((len(c), c.patient_id.nunique()), (3, 2))

    def test_empty_cohort_descriptive_and_export(self):
        p = prepare(small_bundle())
        s = CohortSpec(diagnoses=('NONEXISTENT',))
        c, f = build_cohort(p, s)
        baseline, missingness = descriptive_tables(c)
        self.assertTrue((baseline['可用n'] == 0).all())
        self.assertTrue((missingness['缺失率%'] == 0).all())
        with zipfile.ZipFile(BytesIO(export_package(p, s, c, f))) as z:
            manifest = json.loads(z.read('audit.json'))
            self.assertEqual(manifest['encounters'], 0)
            self.assertEqual(manifest['imputation'], 'none')

    def test_unknown_sex_and_discharge_count_as_missing(self):
        b = small_bundle()
        b.tables['patients'].loc[1, 'sex'] = '未查'
        c, _ = build_cohort(prepare(b), CohortSpec())
        _, m = descriptive_tables(c)
        self.assertEqual(m.set_index('字段').loc['sex', '缺失n'], 1)
        self.assertEqual(m.set_index('字段').loc['discharge_date', '缺失n'], 1)

    def test_invalid_dates_and_chronology_block(self):
        b = small_bundle()
        b.tables['encounters'].loc[0, 'admission_date'] = '01/02/25'
        b.tables['encounters'].loc[1, 'discharge_date'] = '2024-01-01'
        p = prepare(b)
        self.assertEqual(p.usable_encounters.encounter_id.tolist(), ['E3'])
        self.assertIn('DATE_ORDER', p.issues.code.tolist())

    def test_date_only_is_not_an_invented_midnight_baseline(self):
        b = small_bundle([lab('R1', sample_time='2025-01-02')])
        b.tables['encounters'].loc[0, 'admission_date'] = '2025-01-01'
        p = prepare(b)
        c, _ = build_cohort(p, CohortSpec())
        self.assertTrue(c.crp_first_mg_L.isna().all())
        self.assertIn('DATE_ONLY', p.issues.code.tolist())
        exported = csv_bytes(p.usable_encounters).decode('utf-8-sig')
        self.assertNotIn('2025-01-01 00:00:00', exported)

    def test_invalid_cohort_parameters_fail(self):
        p = prepare(small_bundle())
        for spec in [CohortSpec(age_min=90, age_max=20), CohortSpec(baseline_hours=0),
                     CohortSpec(start_date='2026-01-01', end_date='2025-01-01')]:
            with self.assertRaises(ValueError):
                build_cohort(p, spec)


class ImportExportTests(unittest.TestCase):
    def test_csv_and_xlsx_preserve_leading_zero_identifiers(self):
        b = small_bundle()
        csvs = {k: v.to_csv(index=False).encode() for k, v in b.tables.items()}
        self.assertEqual(read_csv_files(csvs).tables['patients'].iloc[0].patient_id, '0001')
        raw = workbook_bytes(b.tables)
        imported = read_workbook(raw)
        self.assertEqual(imported.tables['patients'].iloc[0].patient_id, '0001')

    def test_xlsx_without_dimension_metadata(self):
        b = small_bundle()
        wb = Workbook(write_only=True)
        for name, frame in b.tables.items():
            ws = wb.create_sheet(name)
            ws.append(list(frame.columns))
            for row in frame.itertuples(index=False, name=None):
                ws.append(list(row))
        stream = BytesIO(); wb.save(stream)
        self.assertEqual(len(read_workbook(stream.getvalue()).tables['encounters']), 3)

    def test_formula_cells_rejected_on_import(self):
        wb = load_workbook(BytesIO(workbook_bytes(small_bundle().tables)))
        wb['patients']['B2'] = '=1+1'
        stream = BytesIO(); wb.save(stream)
        with self.assertRaisesRegex(ValueError, '公式'):
            read_workbook(stream.getvalue())

    def test_formula_prefix_escaped_and_numeric_negative_stays_numeric(self):
        frame = pd.DataFrame({'text': ['=1+1', ' +SUM(A1)', '@SUM(A1)', '-cmd'], 'number': [-1, 0, 2, 3]})
        raw = csv_bytes(frame).decode('utf-8-sig')
        self.assertIn("'=1+1", raw)
        wb = load_workbook(BytesIO(workbook_bytes({'test': frame})))
        for row in range(2, 6):
            self.assertNotEqual(wb['test'].cell(row, 1).data_type, 'f')
        self.assertEqual(wb['test']['B2'].value, -1)

    def test_legacy_schema_preserves_pending_identity_and_provenance(self):
        cases = pd.DataFrame([{'原患者ID': '0001', '住院关联键': 'E1', '性别': '女', '年龄（岁）': 65,
                               '入院时间（原试跑表）': '2025-01-01', '出院时间（原试跑表）': '2025-01-10', '原表诊断（待核定）': '待核定原诊断'}])
        labs = pd.DataFrame([{'检验项目关联键': 'R1', '住院关联键': 'E1', '项目名称': 'CRP', '结果原文': '<5',
                              '单位原文': 'mg/L', '标本类型': '血清', '送检时间': '2025-01-02', '报告关联键': 'REP1',
                              '源文件': 'synthetic.json', '源JSON路径': '$.items[0]'}])
        flags = pd.DataFrame([{'核对代码': 'PATIENT_ID_UNVERIFIED'}])
        b = from_legacy_tables({'病例总表': cases, '检验项目明细': labs, '质量核对': flags})
        p = prepare(b)
        self.assertEqual(p.usable_encounters.iloc[0].diagnosis, '待核定原诊断')
        self.assertEqual(p.usable_encounters.iloc[0].procedure, 'UNKNOWN')
        self.assertEqual(p.usable_labs.iloc[0].source_report_id, 'REP1')
        self.assertIn('SOURCE:PATIENT_ID_UNVERIFIED', p.issues.code.tolist())
        self.assertEqual(p.usable_labs.iloc[0].comparator, '<')

    def test_alias_collision_and_missing_required_table_rejected(self):
        b = small_bundle()
        b.tables['patients']['患者编号'] = 'other'
        with self.assertRaisesRegex(ValueError, '重名'):
            from_tables(b.tables)
        with self.assertRaisesRegex(ValueError, '缺少'):
            from_tables({})


if __name__ == '__main__':
    unittest.main()
